from __future__ import annotations

import csv
import json
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl


PROJECT_ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = PROJECT_ROOT / "data_raw"
OUT_DIR = PROJECT_ROOT / "data_processed" / "taas"
DOC_DIR = PROJECT_ROOT / "docs" / "taas"

EXPECTED_ACCIDENT_COUNTS = {
    2020: 8400,
    2021: 8889,
    2022: 9163,
    2023: 8753,
    2024: 9149,
    2025: 8780,
}

TIME_BANDS = [
    "0시~2시",
    "2시~4시",
    "4시~6시",
    "6시~8시",
    "8시~10시",
    "10시~12시",
    "12시~14시",
    "14시~16시",
    "16시~18시",
    "18시~20시",
    "20시~22시",
    "22시~24시",
]

WEEKDAY_SORT = {
    "일": 1,
    "월": 2,
    "화": 3,
    "수": 4,
    "목": 5,
    "금": 6,
    "토": 7,
}

METRIC_MAP = {
    "사고[건]": "사고건수",
    "사망[명]": "사망자수",
    "부상[명]": "부상자수",
}


@dataclass(frozen=True)
class SourceFile:
    path: Path
    table_kind: str
    period: str
    sheet_name: str
    years: list[int]
    max_row: int
    max_column: int


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def to_int(value: Any, *, file_name: str, row: int, col: int, errors: list[dict[str, Any]]) -> int | None:
    if value in (None, ""):
        return None
    try:
        number = int(float(value))
    except (TypeError, ValueError):
        errors.append(
            {
                "file": file_name,
                "row": row,
                "column": col,
                "value": clean_text(value),
                "error": "숫자 변환 오류",
            }
        )
        return None
    return number


def time_start_end(time_band: str) -> tuple[int | None, int | None]:
    match = re.fullmatch(r"(\d+)시~(\d+)시", time_band)
    if not match:
        return None, None
    return int(match.group(1)), int(match.group(2))


def load_sheet(path: Path):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    return workbook.worksheets[0]


def classify_source_file(path: Path) -> SourceFile:
    ws = load_sheet(path)
    headers = [clean_text(ws.cell(1, col).value) for col in range(1, min(ws.max_column, 8) + 1)]
    all_header_text = " ".join(headers)
    if "사고유형" in all_header_text:
        table_kind = "accident_type_time"
    elif "요일" in all_header_text:
        table_kind = "weekday_time"
    else:
        raise ValueError(f"원본 유형을 판단할 수 없습니다: {path.name}")

    years = sorted(
        {
            int(match.group(0))
            for row in (1, 2)
            for col in range(1, ws.max_column + 1)
            for match in re.finditer(r"20\d{2}", clean_text(ws.cell(row, col).value))
        }
    )
    if years == [2020, 2021, 2022, 2023, 2024]:
        period = "2020~2024"
    elif years == [2025]:
        period = "2025"
    else:
        raise ValueError(f"연도 범위를 판단할 수 없습니다: {path.name}, years={years}")

    return SourceFile(
        path=path,
        table_kind=table_kind,
        period=period,
        sheet_name=ws.title,
        years=years,
        max_row=ws.max_row,
        max_column=ws.max_column,
    )


def build_column_map(ws, label_columns: int) -> list[dict[str, Any]]:
    columns: list[dict[str, Any]] = []
    current_year: int | None = None
    for col in range(label_columns + 1, ws.max_column + 1):
        year_text = clean_text(ws.cell(1, col).value)
        if re.fullmatch(r"20\d{2}", year_text):
            current_year = int(year_text)
        elif year_text == "합계":
            current_year = None

        time_band = clean_text(ws.cell(2, col).value)
        if not time_band:
            continue
        columns.append(
            {
                "col": col,
                "year": current_year,
                "time_band": time_band,
                "is_source_total": time_band == "합계" or current_year is None,
            }
        )
    return columns


def add_metric_value(record: dict[str, Any], metric: str, value: int | None) -> None:
    target = METRIC_MAP.get(metric)
    if target is None:
        return
    record[target] = value


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fieldnames})


def parse_weekday_file(source: SourceFile, numeric_errors: list[dict[str, Any]]):
    ws = load_sheet(source.path)
    columns = build_column_map(ws, label_columns=2)
    detail_by_key: dict[tuple[Any, ...], dict[str, Any]] = {}
    total_by_key: dict[tuple[Any, ...], dict[str, Any]] = {}

    current_day = ""
    for row in range(3, ws.max_row + 1):
        first = clean_text(ws.cell(row, 1).value)
        if first.startswith("주"):
            break
        if first:
            current_day = first
        metric = clean_text(ws.cell(row, 2).value)
        if metric not in METRIC_MAP:
            continue

        for col_info in columns:
            col = col_info["col"]
            year = col_info["year"]
            time_band = col_info["time_band"]
            value = to_int(ws.cell(row, col).value, file_name=source.path.name, row=row, col=col, errors=numeric_errors)
            target = total_by_key if col_info["is_source_total"] or current_day == "합계" else detail_by_key
            key = (source.path.name, year, current_day, time_band)
            if key not in target:
                start_hour, end_hour = time_start_end(time_band)
                target[key] = {
                    "원본파일": source.path.name,
                    "연도": year,
                    "요일": current_day,
                    "요일 정렬값": WEEKDAY_SORT.get(current_day, 0),
                    "시간대": time_band,
                    "시작시각": start_hour,
                    "종료시각": end_hour,
                    "원본집계구분": "전체기간합계" if year is None else ("요일합계" if time_band == "합계" else "연도총계"),
                    "사고건수": None,
                    "사망자수": None,
                    "부상자수": None,
                }
            add_metric_value(target[key], metric, value)

    return list(detail_by_key.values()), list(total_by_key.values())


def accident_type_level(major: str, middle: str, detail: str) -> str:
    if major == "합계":
        return "전체총계"
    if middle == "합계":
        return "대분류소계"
    if detail == "합계":
        return "중분류소계"
    if detail:
        return "세부상세"
    if middle:
        return "중분류상세"
    return "대분류단일"


def parse_type_file(source: SourceFile, numeric_errors: list[dict[str, Any]]):
    ws = load_sheet(source.path)
    columns = build_column_map(ws, label_columns=4)
    rows_by_key: dict[tuple[Any, ...], dict[str, Any]] = {}
    totals_by_key: dict[tuple[Any, ...], dict[str, Any]] = {}

    current_major = ""
    current_middle = ""
    active_category = {"major": "", "middle": "", "detail": "", "level": ""}
    for row in range(3, ws.max_row + 1):
        first = clean_text(ws.cell(row, 1).value)
        if first.startswith("주"):
            break
        if first:
            current_major = first
            current_middle = ""
        middle_cell = clean_text(ws.cell(row, 2).value)
        detail_cell = clean_text(ws.cell(row, 3).value)
        if middle_cell:
            current_middle = "" if middle_cell == "합계" else middle_cell

        metric = clean_text(ws.cell(row, 4).value)
        if metric not in METRIC_MAP:
            continue

        # TAAS stores each category as three metric rows. Category labels only
        # appear on the first metric row, so the next two rows must inherit the
        # exact same hierarchy and aggregation level.
        if first or middle_cell or detail_cell:
            major = current_major
            middle = middle_cell if middle_cell else current_middle
            detail = detail_cell
            level = accident_type_level(major, middle_cell, detail)
            active_category = {"major": major, "middle": middle, "detail": detail, "level": level}
        else:
            major = active_category["major"]
            middle = active_category["middle"]
            detail = active_category["detail"]
            level = active_category["level"]

        for col_info in columns:
            col = col_info["col"]
            year = col_info["year"]
            time_band = col_info["time_band"]
            value = to_int(ws.cell(row, col).value, file_name=source.path.name, row=row, col=col, errors=numeric_errors)
            key = (source.path.name, year, major, middle, detail, level, time_band)
            target = totals_by_key if col_info["is_source_total"] or level in {"전체총계", "대분류소계", "중분류소계"} else rows_by_key
            if key not in target:
                start_hour, end_hour = time_start_end(time_band)
                target[key] = {
                    "원본파일": source.path.name,
                    "연도": year,
                    "사고유형 대분류": major,
                    "사고유형 중분류": "" if middle == "합계" else middle,
                    "세부유형": "" if detail == "합계" else detail,
                    "집계수준": level,
                    "시간대": time_band,
                    "시작시각": start_hour,
                    "종료시각": end_hour,
                    "사고건수": None,
                    "사망자수": None,
                    "부상자수": None,
                }
            add_metric_value(target[key], metric, value)

    return list(rows_by_key.values()), list(totals_by_key.values())


def sum_metrics(rows: list[dict[str, Any]], filters: dict[str, Any] | None = None) -> dict[int, dict[str, int]]:
    result: dict[int, dict[str, int]] = defaultdict(lambda: {"사고건수": 0, "사망자수": 0, "부상자수": 0})
    for row in rows:
        if row.get("연도") is None:
            continue
        if filters and any(row.get(key) != value for key, value in filters.items()):
            continue
        year = int(row["연도"])
        for metric in ("사고건수", "사망자수", "부상자수"):
            result[year][metric] += int(row.get(metric) or 0)
    return dict(result)


def rows_for_total(totals: list[dict[str, Any]], *, level: str | None = None, day: str | None = None, time_band: str = "합계"):
    rows = []
    for row in totals:
        if row.get("연도") is None:
            continue
        if row.get("시간대") != time_band:
            continue
        if level is not None and row.get("집계수준") != level:
            continue
        if day is not None and row.get("요일") != day:
            continue
        rows.append(row)
    return rows


def total_map(rows: list[dict[str, Any]]) -> dict[int, dict[str, int]]:
    mapped: dict[int, dict[str, int]] = {}
    for row in rows:
        year = int(row["연도"])
        mapped[year] = {metric: int(row.get(metric) or 0) for metric in ("사고건수", "사망자수", "부상자수")}
    return mapped


def add_validation(
    checks: list[dict[str, Any]],
    name: str,
    status: str,
    detail: str,
    expected: Any = "",
    actual: Any = "",
) -> None:
    checks.append(
        {
            "검증항목": name,
            "상태": status,
            "기대값": expected,
            "실제값": actual,
            "상세": detail,
        }
    )


def compare_total_maps(
    checks: list[dict[str, Any]],
    name: str,
    expected: dict[int, dict[str, int]],
    actual: dict[int, dict[str, int]],
) -> None:
    problems = []
    for year in sorted(set(expected) | set(actual)):
        for metric in ("사고건수", "사망자수", "부상자수"):
            if expected.get(year, {}).get(metric) != actual.get(year, {}).get(metric):
                problems.append(
                    f"{year} {metric}: 기대 {expected.get(year, {}).get(metric)}, 실제 {actual.get(year, {}).get(metric)}"
                )
    add_validation(checks, name, "통과" if not problems else "확인필요", "; ".join(problems) if problems else "일치")


def validation_report(
    weekday_rows: list[dict[str, Any]],
    weekday_totals: list[dict[str, Any]],
    type_rows: list[dict[str, Any]],
    type_totals: list[dict[str, Any]],
    numeric_errors: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    checks: list[dict[str, Any]] = []
    expected_years = set(range(2020, 2026))
    actual_years = {int(row["연도"]) for row in weekday_rows + type_rows if row.get("연도") is not None}
    add_validation(
        checks,
        "2020~2025년 연도 누락 여부",
        "통과" if actual_years == expected_years else "확인필요",
        f"확인 연도: {sorted(actual_years)}",
        sorted(expected_years),
        sorted(actual_years),
    )

    for name, rows, key_fields in [
        ("요일 상세 중복 레코드", weekday_rows, ["연도", "요일", "시간대"]),
        ("사고유형 상세 중복 레코드", type_rows, ["연도", "사고유형 대분류", "사고유형 중분류", "세부유형", "집계수준", "시간대"]),
    ]:
        seen = set()
        duplicates = 0
        missing = 0
        negatives = 0
        for row in rows:
            key = tuple(row.get(field) for field in key_fields)
            duplicates += key in seen
            seen.add(key)
            required_fields = key_fields + ["사고건수", "사망자수", "부상자수"]
            if name.startswith("사고유형"):
                required_fields = [
                    field
                    for field in required_fields
                    if field not in {"사고유형 중분류", "세부유형"}
                ]
                if row.get("집계수준") in {"중분류상세", "세부상세"}:
                    required_fields.append("사고유형 중분류")
                if row.get("집계수준") == "세부상세":
                    required_fields.append("세부유형")
            for field in required_fields:
                missing += row.get(field) in (None, "")
            for metric in ("사고건수", "사망자수", "부상자수"):
                negatives += int(row.get(metric) or 0) < 0
        add_validation(checks, name, "통과" if duplicates == 0 else "확인필요", f"중복 {duplicates}건", 0, duplicates)
        add_validation(checks, f"{name} 결측치", "통과" if missing == 0 else "확인필요", f"결측 {missing}개", 0, missing)
        add_validation(checks, f"{name} 음수", "통과" if negatives == 0 else "확인필요", f"음수 {negatives}개", 0, negatives)

    add_validation(
        checks,
        "숫자 변환 오류",
        "통과" if not numeric_errors else "확인필요",
        f"오류 {len(numeric_errors)}건",
        0,
        len(numeric_errors),
    )

    weekday_total = total_map(rows_for_total(weekday_totals, day="합계"))
    type_total = total_map(rows_for_total(type_totals, level="전체총계"))
    weekday_detail_sum = sum_metrics(weekday_rows)
    type_detail_sum = sum_metrics(type_rows)

    compare_total_maps(checks, "요일 시간대 합계와 원본 총계 일치 여부", weekday_total, weekday_detail_sum)
    compare_total_maps(checks, "사고유형 상세 합계와 연도별 총계 일치 여부", type_total, type_detail_sum)
    compare_total_maps(checks, "두 통계의 연도별 사고·사망·부상 총계 일치 여부", weekday_total, type_total)

    weekday_accident_reference = {year: values["사고건수"] for year, values in weekday_total.items()}
    reference_problems = [
        f"{year}: 참고 {EXPECTED_ACCIDENT_COUNTS.get(year)}, 원본 {weekday_accident_reference.get(year)}"
        for year in sorted(EXPECTED_ACCIDENT_COUNTS)
        if EXPECTED_ACCIDENT_COUNTS.get(year) != weekday_accident_reference.get(year)
    ]
    add_validation(
        checks,
        "검증 참고용 연도별 사고건수와 원본 총계 비교",
        "통과" if not reference_problems else "확인필요",
        "; ".join(reference_problems) if reference_problems else "일치",
    )

    weekday_day_sum = defaultdict(lambda: {"사고건수": 0, "사망자수": 0, "부상자수": 0})
    for row in weekday_totals:
        if row.get("연도") is not None and row.get("시간대") == "합계" and row.get("요일") not in ("합계", ""):
            year = int(row["연도"])
            for metric in ("사고건수", "사망자수", "부상자수"):
                weekday_day_sum[year][metric] += int(row.get(metric) or 0)
    compare_total_maps(checks, "요일 합계와 연도별 총계 일치 여부", weekday_total, dict(weekday_day_sum))

    major_sum = defaultdict(lambda: {"사고건수": 0, "사망자수": 0, "부상자수": 0})
    for row in type_totals:
        if row.get("연도") is not None and row.get("시간대") == "합계" and row.get("집계수준") == "대분류소계":
            year = int(row["연도"])
            for metric in ("사고건수", "사망자수", "부상자수"):
                major_sum[year][metric] += int(row.get(metric) or 0)
    for row in type_totals:
        if row.get("연도") is not None and row.get("시간대") == "합계" and row.get("집계수준") == "대분류단일":
            year = int(row["연도"])
            for metric in ("사고건수", "사망자수", "부상자수"):
                major_sum[year][metric] += int(row.get(metric) or 0)
    compare_total_maps(checks, "사고유형 대분류 합계와 연도별 총계 일치 여부", type_total, dict(major_sum))

    return checks


def hierarchy_summary(type_rows: list[dict[str, Any]], type_totals: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    all_rows = [row for row in type_rows + type_totals if row.get("연도") is not None and row.get("시간대") == "합계"]
    by_year = defaultdict(lambda: defaultdict(int))
    for row in all_rows:
        by_year[int(row["연도"])][row["집계수준"]] += 1
    for year in sorted(by_year):
        item = {"연도": year}
        item.update(dict(sorted(by_year[year].items())))
        rows.append(item)
    return rows


def render_markdown(
    sources: list[SourceFile],
    checks: list[dict[str, Any]],
    hierarchy_rows: list[dict[str, Any]],
    row_counts: dict[str, int],
) -> str:
    source_lines = "\n".join(
        f"- `{source.path.name}`: {source.table_kind}, {source.period}, 시트 `{source.sheet_name}`, "
        f"{source.max_row}행 x {source.max_column}열"
        for source in sources
    )
    check_lines = "\n".join(
        f"- {row['검증항목']}: {row['상태']}. {row['상세']}" for row in checks
    )
    hierarchy_lines = "\n".join(
        "- "
        + ", ".join(f"{key}={value}" for key, value in row.items())
        for row in hierarchy_rows
    )
    count_lines = "\n".join(f"- `{name}`: {count}행" for name, count in row_counts.items())

    return f"""# TAAS 어린이 교통사고 전국 집계 통계 전처리 및 검증

생성일: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}

## 원본 파일과 데이터 범위

{source_lines}

자료는 12세 이하 어린이 사상자가 발생한 전국 어린이 교통사고 집계다. 지역명, 학교명, 좌표가 없는 전국 단위 통계이므로 시군구별 예측 결과에 같은 값을 복제해 붙이지 않는다.

## 정제 CSV

{count_lines}

### `taas_weekday_time_2020_2025.csv`

- `연도`: 집계 연도
- `요일`: 일, 월, 화, 수, 목, 금, 토
- `요일 정렬값`: 일요일 1부터 토요일 7까지의 정렬용 값
- `시간대`: 2시간 단위 시간대
- `시작시각`, `종료시각`: 시간대 경계의 24시간 정수
- `사고건수`: 사고 발생 건수
- `사망자수`: 사망자 수
- `부상자수`: 부상자 수

### `taas_accident_type_time_2020_2025.csv`

- `연도`: 집계 연도
- `사고유형 대분류`, `사고유형 중분류`, `세부유형`: 원본 사고유형 계층
- `집계수준`: `중분류상세`, `세부상세`, `대분류단일` 등 원본 행의 집계 수준
- `시간대`, `시작시각`, `종료시각`: 2시간 단위 시간대와 경계
- `사고건수`, `사망자수`, `부상자수`: 원본 수치. 사고건수와 사상자수는 서로 다른 단위다.

총계와 소계는 `taas_weekday_time_totals_2020_2025.csv`, `taas_accident_type_time_totals_2020_2025.csv`에 별도 저장했다. 분석 시 총계, 대분류소계, 중분류소계를 상세 행과 함께 합산하면 중복 집계된다.

## 전처리 과정

1. `data_raw`의 엑셀 4개를 내부 헤더 기준으로 요일별·시간대별 2020~2024, 요일별·시간대별 2025, 사고유형별·시간대별 2020~2024, 사고유형별·시간대별 2025로 분류했다.
2. 병합 셀 때문에 비어 있는 연도, 요일, 사고유형 계층 값을 직전 유효 값으로 보정했다.
3. 2020~2024 파일의 전체기간 합계 열과 각 연도의 합계 열은 검증용 totals CSV에 분리했다.
4. 시간대 상세 열만 분석용 CSV에 펼쳤다.
5. 사고유형 계층은 원본에 없는 유형을 만들지 않고, 2025년에 없는 `도로이탈 기타` 같은 항목을 0으로 채우지 않았다.

## 사고유형 계층 차이

{hierarchy_lines}

2020~2024년에는 차량단독 `도로이탈` 아래에 `도로이탈 추락`, `도로이탈 기타` 세부유형과 `도로이탈` 중분류소계가 있다. 2025년에는 `도로이탈 추락`만 세부유형으로 제공되고 별도 `도로이탈 기타` 또는 `도로이탈` 소계 행이 없다.

## 검증 결과

{check_lines}

## 서비스에서 활용할 수 있는 차트

- 연도별 시간대 사고건수 라인 또는 히트맵
- 요일 x 시간대 사고건수 히트맵
- 요일별 사망자수·부상자수 막대그래프
- 사고유형 대분류별 연도 추이
- 사고유형 중분류 또는 세부유형별 시간대 분포

## 자료의 한계

- 전국 집계 통계라 지역별 위험도를 직접 설명하지 못한다.
- 지역명, 학교명, 좌표가 없어 지도 기반 위험지역 분석에는 사용할 수 없다.
- 어린이 전체 교통사고 통계이며 보행 어린이 사고만을 의미하지 않는다.
- 사고건수와 사망자수·부상자수는 단위가 다르므로 같은 의미로 비교하거나 합산하지 않는다.
- 2025년 사고유형 계층 일부가 2020~2024년과 다르므로 유형별 장기 비교 때 원본 계층 차이를 함께 표시해야 한다.
"""


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    DOC_DIR.mkdir(parents=True, exist_ok=True)

    raw_files = sorted(RAW_DIR.glob("*.xls*"))
    if len(raw_files) != 4:
        raise RuntimeError(f"data_raw 엑셀 4개가 필요합니다. 현재 {len(raw_files)}개")

    sources = [classify_source_file(path) for path in raw_files]
    numeric_errors: list[dict[str, Any]] = []
    weekday_rows: list[dict[str, Any]] = []
    weekday_totals: list[dict[str, Any]] = []
    type_rows: list[dict[str, Any]] = []
    type_totals: list[dict[str, Any]] = []

    for source in sources:
        if source.table_kind == "weekday_time":
            detail, totals = parse_weekday_file(source, numeric_errors)
            weekday_rows.extend(detail)
            weekday_totals.extend(totals)
        elif source.table_kind == "accident_type_time":
            detail, totals = parse_type_file(source, numeric_errors)
            type_rows.extend(detail)
            type_totals.extend(totals)

    weekday_rows.sort(key=lambda row: (row["연도"], row["요일 정렬값"], row["시작시각"]))
    weekday_totals.sort(key=lambda row: (row["원본파일"], row.get("연도") or 0, row.get("요일", ""), row.get("시작시각") or -1))
    type_rows.sort(
        key=lambda row: (
            row["연도"],
            row["사고유형 대분류"],
            row["사고유형 중분류"],
            row["세부유형"],
            row["집계수준"],
            row["시작시각"],
        )
    )
    type_totals.sort(
        key=lambda row: (
            row["원본파일"],
            row.get("연도") or 0,
            row.get("사고유형 대분류", ""),
            row.get("사고유형 중분류", ""),
            row.get("세부유형", ""),
            row.get("집계수준", ""),
            row.get("시작시각") or -1,
        )
    )

    weekday_fields = ["연도", "요일", "요일 정렬값", "시간대", "시작시각", "종료시각", "사고건수", "사망자수", "부상자수"]
    weekday_total_fields = ["원본파일", "연도", "요일", "요일 정렬값", "시간대", "시작시각", "종료시각", "원본집계구분", "사고건수", "사망자수", "부상자수"]
    type_fields = ["연도", "사고유형 대분류", "사고유형 중분류", "세부유형", "집계수준", "시간대", "시작시각", "종료시각", "사고건수", "사망자수", "부상자수"]
    type_total_fields = ["원본파일"] + type_fields

    write_csv(OUT_DIR / "taas_weekday_time_2020_2025.csv", weekday_rows, weekday_fields)
    write_csv(OUT_DIR / "taas_weekday_time_totals_2020_2025.csv", weekday_totals, weekday_total_fields)
    write_csv(OUT_DIR / "taas_accident_type_time_2020_2025.csv", type_rows, type_fields)
    write_csv(OUT_DIR / "taas_accident_type_time_totals_2020_2025.csv", type_totals, type_total_fields)

    source_inventory = [
        {
            "원본파일": source.path.name,
            "자료구분": source.table_kind,
            "기간": source.period,
            "시트": source.sheet_name,
            "연도": ",".join(str(year) for year in source.years),
            "행수": source.max_row,
            "열수": source.max_column,
        }
        for source in sources
    ]
    write_csv(
        OUT_DIR / "taas_source_inventory.csv",
        source_inventory,
        ["원본파일", "자료구분", "기간", "시트", "연도", "행수", "열수"],
    )

    hierarchy_rows = hierarchy_summary(type_rows, type_totals)
    hierarchy_fields = sorted({key for row in hierarchy_rows for key in row}, key=lambda value: ("연도" != value, value))
    write_csv(OUT_DIR / "taas_accident_type_hierarchy_summary.csv", hierarchy_rows, hierarchy_fields)

    checks = validation_report(weekday_rows, weekday_totals, type_rows, type_totals, numeric_errors)
    write_csv(OUT_DIR / "taas_validation_summary.csv", checks, ["검증항목", "상태", "기대값", "실제값", "상세"])
    (OUT_DIR / "taas_validation_summary.json").write_text(
        json.dumps({"checks": checks, "numeric_errors": numeric_errors}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    row_counts = {
        "taas_weekday_time_2020_2025.csv": len(weekday_rows),
        "taas_weekday_time_totals_2020_2025.csv": len(weekday_totals),
        "taas_accident_type_time_2020_2025.csv": len(type_rows),
        "taas_accident_type_time_totals_2020_2025.csv": len(type_totals),
        "taas_source_inventory.csv": len(source_inventory),
        "taas_accident_type_hierarchy_summary.csv": len(hierarchy_rows),
        "taas_validation_summary.csv": len(checks),
    }
    markdown = render_markdown(sources, checks, hierarchy_rows, row_counts)
    (DOC_DIR / "taas_preprocessing_validation.md").write_text(markdown, encoding="utf-8")

    print(json.dumps({"row_counts": row_counts, "checks": checks}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
